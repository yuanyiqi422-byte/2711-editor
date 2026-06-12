"""从 TXT / DOCX / PDF / EPUB 等文件提取正文，支持大文件与扫描版 PDF OCR。"""

from __future__ import annotations

import io
import os
import re
import tempfile
from dataclasses import dataclass
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

def _env_int(name: str, default: str, *, zero_unlimited: bool = False) -> int:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return int(default)
    try:
        value = int(str(raw).strip())
    except ValueError:
        print(f"[2711编辑器] 警告: {name}={raw!r} 无效，使用默认 {default}")
        return int(default)
    if zero_unlimited and value == 0:
        return 0
    return value


_max_upload_mb = _env_int("MAX_UPLOAD_MB", "0", zero_unlimited=True)
MAX_UPLOAD_BYTES = _max_upload_mb * 1024 * 1024 if _max_upload_mb > 0 else 0


def upload_max_mb() -> int:
    """Return max upload size in MB for health API (0 = unlimited)."""
    if MAX_UPLOAD_BYTES <= 0:
        return 0
    return MAX_UPLOAD_BYTES // (1024 * 1024)


# 0 = 不限制
MAX_EXTRACT_CHARS = _env_int("MAX_EXTRACT_CHARS", "0", zero_unlimited=True)
MAX_CREATE_CHARS = _env_int("MAX_CREATE_CHARS", "0", zero_unlimited=True)
# 0 = 全文送入 LLM，不智能采样；>0 为显式采样上限
LLM_INPUT_CHARS = _env_int("LLM_INPUT_CHARS", "0", zero_unlimited=True)
LLM_MODEL_SAFE_CHARS = _env_int("LLM_MODEL_SAFE_CHARS", "0", zero_unlimited=True)


def effective_llm_limit() -> int:
    """送入 LLM 的字数上限；0 表示不截取。"""
    if LLM_INPUT_CHARS > 0:
        return LLM_INPUT_CHARS
    return LLM_MODEL_SAFE_CHARS


# 0 = 处理全部 PDF 页，无页数上限
MAX_PDF_PAGES = _env_int("MAX_PDF_PAGES", "0", zero_unlimited=True)
PDF_PAGE_HARD_CAP = _env_int("PDF_PAGE_HARD_CAP", "0", zero_unlimited=True)
OCR_MIN_CHARS_PER_PAGE = 40
_CHAPTER_HEADING = re.compile(
    r"^\s*(?:第[零一二三四五六七八九十百千万0-9]+[章节回幕]|Chapter\s+\d+)",
    re.MULTILINE | re.IGNORECASE,
)

ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".docx",
    ".pdf",
    ".rtf",
    ".epub",
}
EXTENSION_LABELS = {
    ".txt": "纯文本",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".docx": "Word 文档",
    ".pdf": "PDF",
    ".rtf": "RTF",
    ".epub": "EPUB",
}


class ExtractError(Exception):
    pass


@dataclass
class ExtractResult:
    text: str
    char_count: int
    truncated: bool
    truncate_note: str
    format: str
    filename: str
    llm_text: str = ""
    llm_char_count: int = 0
    pages_processed: int | None = None
    total_pages: int | None = None
    used_ocr: bool = False
    warnings: list[str] | None = None


def validate_upload(filename: str, size: int) -> str:
    if size <= 0:
        raise ExtractError("文件为空。")
    if MAX_UPLOAD_BYTES > 0 and size > MAX_UPLOAD_BYTES:
        mb = MAX_UPLOAD_BYTES // (1024 * 1024)
        raise ExtractError(f"文件超过 {mb} MB 上限，请拆分章节后分别上传。")
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        allowed = "、".join(sorted(ALLOWED_EXTENSIONS))
        raise ExtractError(f"不支持「{ext or '未知'}」格式。当前支持：{allowed}")
    if ext == ".doc":
        raise ExtractError("旧版 .doc 请先在 Word 中另存为 .docx 再上传。")
    return ext


def _normalize_text(raw: str) -> str:
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return text.strip()


def _find_chapter_spans(text: str) -> list[tuple[int, int, str]]:
    matches = list(_CHAPTER_HEADING.finditer(text))
    if len(matches) < 4:
        return []
    spans: list[tuple[int, int, str]] = []
    for i, m in enumerate(matches):
        start = m.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        spans.append((start, end, m.group(0).strip()[:48]))
    return spans


def segment_text_for_slots(text: str, slot_count: int) -> list[str]:
    """将正文切成 slot_count 段；优先按章节均匀分组，否则按字数均分。"""
    body = text.strip()
    if not body or slot_count <= 0:
        return []
    if slot_count == 1:
        return [body]

    spans = _find_chapter_spans(body)
    if len(spans) >= slot_count:
        chapter_count = len(spans)
        groups: list[list[int]] = [[] for _ in range(slot_count)]
        for i in range(chapter_count):
            group_idx = min(int(i * slot_count / chapter_count), slot_count - 1)
            groups[group_idx].append(i)
        parts: list[str] = []
        for group in groups:
            if not group:
                parts.append("")
                continue
            start = spans[group[0]][0]
            end = spans[group[-1]][1]
            parts.append(body[start:end].strip())
        return parts

    size = max(1, len(body) // slot_count)
    parts: list[str] = []
    for i in range(slot_count):
        start = i * size
        end = len(body) if i == slot_count - 1 else (i + 1) * size
        parts.append(body[start:end].strip())
    return parts


def _truncate_by_chapters(text: str, max_chars: int) -> tuple[str, bool, str] | None:
    spans = _find_chapter_spans(text)
    if len(spans) < 4:
        return None

    n = len(spans)
    avg = max(400, len(text) // n)
    pick_count = max(4, min(n, max_chars // avg))
    if pick_count >= n and len(text) <= max_chars:
        return None

    if pick_count >= n:
        pick_count = max(4, n // 2)

    if pick_count < 4:
        return None

    indices = (
        list(range(pick_count))
        if pick_count >= n
        else [int(i * (n - 1) / (pick_count - 1)) for i in range(pick_count)]
    )

    parts: list[str] = []
    used = 0
    sep = "\n\n⋯⋯（章节省略）⋯⋯\n\n"
    sep_len = len(sep)

    for j, idx in enumerate(indices):
        start, end, title = spans[idx]
        chunk = text[start:end].strip()
        if not chunk:
            continue
        budget_left = max_chars - used - (sep_len if parts else 0)
        if budget_left < 400:
            break
        if len(chunk) > budget_left:
            chunk = chunk[:budget_left]
        header = f"\n\n【{title}】\n\n" if title else "\n\n"
        piece = header + chunk
        if used + len(piece) > max_chars:
            piece = piece[: max_chars - used]
        parts.append(piece)
        used += len(piece)
        if j + 1 < len(indices) and used + sep_len < max_chars:
            parts.append(sep)
            used += sep_len

    body = "".join(parts)[:max_chars].strip()
    if len(body) < max_chars // 4:
        return None

    skipped = n - len(indices)
    note = (
        f"原文共 {len(text):,} 字、识别约 {n} 个章节，"
        f"已均匀采样 {len(indices)} 章（约 {len(body):,} 字）用于生成游戏。"
    )
    if skipped:
        note += f" 其余 {skipped} 章未载入，可按章拆分上传或粘贴关键章节。"
    return body, True, note


def prepare_llm_text(text: str, max_chars: int | None = None) -> tuple[str, bool, str]:
    """为 LLM 准备正文；过长时智能采样，不拒绝用户全文提交。"""
    limit = effective_llm_limit() if max_chars is None else max_chars
    if limit <= 0:
        return text, False, ""
    return smart_truncate(text, limit)


def smart_truncate(text: str, max_chars: int) -> tuple[str, bool, str]:
    if max_chars <= 0 or len(text) <= max_chars:
        return text, False, ""

    by_chapter = _truncate_by_chapters(text, max_chars)
    if by_chapter is not None:
        return by_chapter

    head = int(max_chars * 0.42)
    mid_len = int(max_chars * 0.16)
    tail = max_chars - head - mid_len - 120
    mid_start = max(0, (len(text) - mid_len) // 2)
    mid = text[mid_start : mid_start + mid_len]
    note = (
        f"原文共 {len(text):,} 字，已智能截取开头、中段与结尾共 {max_chars:,} 字用于生成游戏。"
        "如需完整覆盖，建议按章节拆分上传。"
    )
    chunk = (
        text[:head]
        + "\n\n⋯⋯（中段省略，原文过长）⋯⋯\n\n"
        + mid
        + "\n\n⋯⋯（后段接续）⋯⋯\n\n"
        + text[-tail:]
    )
    return chunk[:max_chars], True, note


def _effective_pdf_page_limit(max_pages: int) -> int:
    """0 表示不限制页数。"""
    if max_pages <= 0:
        if PDF_PAGE_HARD_CAP <= 0:
            return 0
        return PDF_PAGE_HARD_CAP
    if PDF_PAGE_HARD_CAP <= 0:
        return max_pages
    return min(max_pages, PDF_PAGE_HARD_CAP)


def _select_pdf_page_indices(page_count: int, max_pages: int) -> tuple[list[int], str]:
    limit = _effective_pdf_page_limit(max_pages)
    if limit <= 0 or page_count <= limit:
        return list(range(page_count)), ""

    head_n = max(1, int(limit * 0.5))
    mid_n = max(1, int(limit * 0.25))
    tail_n = max(1, limit - head_n - mid_n)

    head = list(range(head_n))
    mid_start = max(head_n, (page_count - mid_n) // 2)
    mid = list(range(mid_start, min(mid_start + mid_n, page_count)))
    tail_start = max(0, page_count - tail_n)
    tail = list(range(tail_start, page_count))

    indices = sorted(set(head + mid + tail))
    if max_pages <= 0:
        note = (
            f"PDF 共 {page_count} 页，超过安全上限 {PDF_PAGE_HARD_CAP} 页，"
            f"已采样 {len(indices)} 页（开头+中段+结尾）。"
        )
    else:
        note = (
            f"PDF 共 {page_count} 页，已采样 {len(indices)} 页"
            f"（开头 {len(head)} + 中段 {len(mid)} + 结尾 {len(tail)}）。"
            f"可在 backend/.env 设置 MAX_PDF_PAGES=0 处理全部页面。"
        )
    return indices, note


def _page_text_from_fitz(page) -> str:
    try:
        blocks = page.get_text("blocks")
    except Exception:
        return page.get_text("text").strip()

    typed: list[tuple[float, float, str]] = []
    for block in blocks:
        if len(block) < 5:
            continue
        x0, y0, _x1, _y1, text = block[0], block[1], block[2], block[3], block[4]
        t = (text or "").strip()
        if t:
            typed.append((y0, x0, t))

    if not typed:
        return page.get_text("text").strip()

    typed.sort(key=lambda item: (round(item[0], 1), item[1]))
    return "\n".join(t for _, _, t in typed)


def _strip_repeated_margin_lines(page_texts: list[str]) -> list[str]:
    if len(page_texts) < 4:
        return page_texts

    top: dict[str, int] = {}
    bottom: dict[str, int] = {}
    for raw in page_texts:
        lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        if not lines:
            continue
        if len(lines[0]) <= 80:
            top[lines[0]] = top.get(lines[0], 0) + 1
        if len(lines[-1]) <= 80:
            bottom[lines[-1]] = bottom.get(lines[-1], 0) + 1

    threshold = max(3, len(page_texts) // 3)
    drop_top = {k for k, v in top.items() if v >= threshold and len(k) >= 2}
    drop_bottom = {k for k, v in bottom.items() if v >= threshold and len(k) >= 2}

    cleaned: list[str] = []
    for raw in page_texts:
        lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        if not lines:
            cleaned.append(raw)
            continue
        while lines and lines[0] in drop_top:
            lines.pop(0)
        while lines and lines[-1] in drop_bottom:
            lines.pop()
        cleaned.append("\n".join(lines))
    return cleaned


def _decode_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractError("无法识别文本编码，请另存为 UTF-8 的 .txt 后重试。")


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractError("服务器未安装 python-docx，无法解析 Word。") from exc

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if not parts:
        raise ExtractError("Word 文档中未检测到正文，请确认不是空白文档。")
    return "\n\n".join(parts)


def _extract_rtf(data: bytes) -> str:
    raw = _decode_txt(data)
    text = re.sub(r"\\[a-z]+\d*\s?", "", raw)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\\'[0-9a-f]{2}", "", text, flags=re.I)
    if len(text.strip()) < 80:
        raise ExtractError("RTF 解析结果过短，建议另存为 .docx 或 .txt。")
    return text


def _pdf_needs_ocr(page_texts: list[str]) -> bool:
    if not page_texts:
        return True
    nonempty = [t for t in page_texts if len(t.strip()) >= OCR_MIN_CHARS_PER_PAGE]
    return len(nonempty) < max(1, len(page_texts) // 4)


def _ocr_pdf_pages(doc, page_indices: list[int], warnings: list[str]) -> list[str]:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise ExtractError(
            "该 PDF 为扫描版，需要 OCR 组件。请在后端执行：pip install rapidocr-onnxruntime"
        ) from exc

    ocr = RapidOCR()
    scale = 2.0
    out: list[str] = []
    warnings.append("已启用 OCR 识别扫描版 PDF，大文件可能需 1～3 分钟。")
    for idx in page_indices:
        page = doc.load_page(idx)
        import fitz as _fitz

        pix = page.get_pixmap(matrix=_fitz.Matrix(scale, scale))
        result, _ = ocr(pix.tobytes("png"))
        if not result:
            out.append("")
            continue
        out.append("\n".join(line[1] for line in result if line[1].strip()))
    return out


def _extract_pdf(data: bytes, warnings: list[str]) -> tuple[str, int, bool, int]:
    try:
        import fitz
    except ImportError as exc:
        raise ExtractError("服务器未安装 pymupdf，无法解析 PDF。请 pip install pymupdf") from exc

    doc = fitz.open(stream=data, filetype="pdf")
    page_count = doc.page_count
    if page_count == 0:
        raise ExtractError("PDF 没有页面。")

    indices, sample_note = _select_pdf_page_indices(page_count, MAX_PDF_PAGES)
    if sample_note:
        warnings.append(sample_note)

    page_texts = [_page_text_from_fitz(doc.load_page(i)) for i in indices]
    page_texts = _strip_repeated_margin_lines(page_texts)
    used_ocr = False

    if _pdf_needs_ocr(page_texts):
        used_ocr = True
        ocr_note = (
            f"扫描版 PDF：对采样的 {len(indices)} 页执行 OCR"
            + (f"（全书 {page_count} 页）。" if page_count > len(indices) else "。")
        )
        warnings.append(ocr_note)
        page_texts = _ocr_pdf_pages(doc, indices, warnings)

    doc.close()
    body = "\n\n".join(t for t in page_texts if t.strip())
    if len(body.strip()) < 80:
        raise ExtractError(
            "PDF 未能提取足够正文。若为扫描件，请确认已安装 pymupdf 与 rapidocr-onnxruntime；"
            "也可尝试将 PDF 按章导出为 Word/TXT 后上传。"
        )
    return body, len(indices), used_ocr, page_count


class _HTMLTextExtractor(HTMLParser):
    _BLOCK = frozenset(
        {"p", "div", "br", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "section", "article"}
    )
    _SKIP = frozenset({"script", "style", "head", "meta", "link", "svg", "noscript"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "br":
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCK:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth or not data.strip():
            return
        self._chunks.append(data)

    def get_text(self) -> str:
        return unescape("".join(self._chunks))


def _html_to_text(html: str) -> str:
    parser = _HTMLTextExtractor()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        text = re.sub(r"<[^>]+>", " ", html)
        return unescape(re.sub(r"\s+", " ", text).strip())
    text = parser.get_text()
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _epub_chapter_text(item) -> str:
    raw = item.get_content()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            html = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        html = raw.decode("utf-8", errors="replace")
    return _html_to_text(html)


def _extract_epub(data: bytes, warnings: list[str]) -> str:
    try:
        import ebooklib
        from ebooklib import epub
    except ImportError as exc:
        raise ExtractError("服务器未安装 EbookLib，无法解析 EPUB。请 pip install EbookLib") from exc

    try:
        book = epub.read_epub(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError(f"无法打开 EPUB 文件：{exc}") from exc

    spine_ids: list[str] = []
    for entry in book.spine:
        idref = entry[0] if isinstance(entry, (list, tuple)) else entry
        if idref:
            spine_ids.append(idref)

    parts: list[str] = []
    skipped = 0

    def append_item(item) -> None:
        nonlocal skipped
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT:
            return
        try:
            chapter = _epub_chapter_text(item)
        except Exception:
            skipped += 1
            return
        if chapter:
            parts.append(chapter)

    for idref in spine_ids:
        append_item(book.get_item_with_id(idref))

    if not parts:
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            append_item(item)

    if skipped:
        warnings.append(f"EPUB 中有 {skipped} 个章节解析失败，已跳过。")
    if not parts:
        raise ExtractError("EPUB 中未检测到正文，请确认文件包含可读章节。")

    body = "\n\n".join(parts)
    if len(body.strip()) < 40:
        raise ExtractError("EPUB 提取的正文过短，请确认不是空白或加密电子书。")
    return body


def extract_from_bytes(data: bytes, filename: str) -> ExtractResult:
    ext = validate_upload(filename, len(data))
    warnings: list[str] = []

    total_pages: int | None = None

    if ext in {".txt", ".md", ".markdown"}:
        raw = _decode_txt(data)
        fmt = EXTENSION_LABELS.get(ext, "文本")
        pages = None
        ocr = False
    elif ext == ".docx":
        raw = _extract_docx(data)
        fmt = "Word"
        pages = None
        ocr = False
    elif ext == ".rtf":
        raw = _extract_rtf(data)
        fmt = "RTF"
        pages = None
        ocr = False
    elif ext == ".pdf":
        raw, pages, ocr, total_pages = _extract_pdf(data, warnings)
        fmt = "PDF"
    elif ext == ".epub":
        raw = _extract_epub(data, warnings)
        fmt = "EPUB"
        pages = None
        ocr = False
    else:
        raise ExtractError(f"未实现的格式：{ext}")

    text = _normalize_text(raw)
    if MAX_EXTRACT_CHARS > 0 and len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS]
        warnings.append(f"提取结果已限制在 {MAX_EXTRACT_CHARS:,} 字以内。")

    llm_body, llm_truncated, llm_note = prepare_llm_text(text)

    return ExtractResult(
        text=text,
        char_count=len(text),
        truncated=llm_truncated,
        truncate_note=llm_note,
        llm_text=llm_body,
        llm_char_count=len(llm_body),
        format=fmt,
        filename=Path(filename).name,
        pages_processed=pages,
        total_pages=total_pages if ext == ".pdf" else None,
        used_ocr=ocr if ext == ".pdf" else False,
        warnings=warnings or None,
    )


def extract_from_path(path: Path) -> ExtractResult:
    return extract_from_bytes(path.read_bytes(), path.name)
